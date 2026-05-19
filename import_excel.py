import openpyxl
import os
import sys
import re

# Thêm thư mục hiện tại vào path để import được src
sys.path.append(os.getcwd())

from src.core.db_manager import RoKDatabase

def parse_combined_row(raw_text):
    """Parse câu hỏi dạng gộp câu hỏi + câu trả lời từ cột A khi cột B trống"""
    cleaned = raw_text.strip()
    # Loại bỏ số thứ tự ở đầu ví dụ: 1., 166.
    cleaned = re.sub(r'^\d+[\.\s\-]+', '', cleaned).strip()
    
    parts = []
    # Nếu có dấu hỏi chấm, ưu tiên cắt theo dấu hỏi chấm
    if '?' in cleaned:
        idx = cleaned.find('?')
        p1 = cleaned[:idx+1].strip()
        p2 = cleaned[idx+1:].strip()
        # Loại bỏ các ký tự phân tách thừa ở đầu phần 2
        p2 = re.sub(r'^[\s\:\/\-\.\,\"]+', '', p2).strip()
        if p1 and p2:
            parts = [p1, p2]
            
    # Nếu không có dấu hỏi chấm hoặc không cắt được, cắt theo các ký tự phân tách thông thường
    if not parts:
        for sep in ['/.', '/', ':', ' - ', ' – ']:
            if sep in cleaned:
                parts = [p.strip() for p in cleaned.split(sep, 1)]
                break
                
    if len(parts) == 2:
        p1, p2 = parts[0], parts[1]
        
        # Nhận diện phần nào là Câu hỏi và phần nào là Đáp án
        if '?' in p1:
            q, a = p1, p2
        elif '?' in p2:
            q, a = p2, p1
        else:
            # Quy tắc: Phần dài hơn là Câu hỏi, phần ngắn hơn là Đáp án
            if len(p1) > len(p2):
                q, a = p1, p2
            else:
                q, a = p2, p1
                
        # Làm sạch ký tự phân tách thừa
        q = re.sub(r'^[\s\:\/\-\.\,\"]+', '', q).strip()
        a = re.sub(r'^[\s\:\/\-\.\,\"]+', '', a).strip()
        a = re.sub(r'[\.\"]+$', '', a).strip()
        
        if len(q) > 8 and len(a) > 0:
            return q, a
            
    return None, None

def import_from_excel(file_path):
    if not os.path.exists(file_path):
        print(f"Lỗi: Không tìm thấy file {file_path}")
        return

    print(f"--- Đang đọc file: {file_path} ---")
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb['Sheet2'] if 'Sheet2' in wb.sheetnames else wb.active
        print(f"Đang xử lý sheet: {sheet.title}")

        db = RoKDatabase()
        
        count = 0
        duplicate = 0
        skipped = 0
        
        for row in sheet.iter_rows(min_row=1, values_only=True):
            if not row or not any(row):
                continue
                
            question = str(row[0]).strip() if row[0] else ""
            answer = str(row[1]).strip() if row[1] else ""
            
            # Nếu Column B rỗng, thử parse từ Column A dạng gộp
            if question and not answer:
                parsed_q, parsed_a = parse_combined_row(question)
                if parsed_q and parsed_a:
                    question, answer = parsed_q, parsed_a
                else:
                    skipped += 1
                    continue
            
            # Bỏ qua nếu dữ liệu trống hoặc tiêu đề
            if not question or not answer or question.lower() == "câu hỏi":
                continue
            
            # Mặc định thêm vào game Rise of Kingdoms
            if db.add_question("Rise of Kingdoms", question, answer):
                count += 1
            else:
                duplicate += 1
                
        print(f"Hoàn tất! Thêm mới: {count} câu. Trùng lặp: {duplicate} câu. Bỏ qua (không parse được): {skipped} câu.")
        
    except Exception as e:
        print(f"Lỗi khi xử lý: {e}")

if __name__ == "__main__":
    file_name = "ROK Question.xlsx"
    import_from_excel(file_name)
